import json, sys
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

src, out, orig = sys.argv[1], sys.argv[2], sys.argv[3]
d = json.load(open(src))
O = json.load(open(orig))
_A = json.load(open(sys.argv[4]))
ADD, REMAP = _A['add'], _A['remap']
import re
_oid = lambda o: (re.search(r'album/(\w{22})', o['url'] or '') or [None, ''])[1]
oid = lambda o: REMAP.get(_oid(o), _oid(o))
NEW = PatternFill('solid', fgColor='1C5CAB')
TYPES = {'album': 'Álbum', 'compilation': 'Recopilatorio', 'single': 'Sencillo o EP'}
F = Font(name='Arial', size=10)
H = Font(name='Arial', size=10, bold=True, color='FFFFFF')
FILL = PatternFill('solid', fgColor='1F2937')
LINK = Font(name='Arial', size=10, color='0563C1', underline='single')

def sheet(ws, headers, widths, rows, notes):
    ws.append(headers)
    for c, w in zip(ws[1], widths):
        c.font, c.fill = H, FILL
        c.alignment = Alignment(vertical='center')
        ws.column_dimensions[c.column_letter].width = w
    for col, text in notes.items():
        ws[f'{col}1'].comment = Comment(text, 'Hill-B')
    for r in rows:
        ws.append(r)
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{ws.max_row}'
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = F

wb = Workbook()
ws = wb.active
ws.title = 'Canciones'
rows = []
for i, s in enumerate(d['songs'], start=2):
    o = O[s['id']]
    rows.append([s['pos'], s['name'], s['artists'], s['album'], s['albumArtists'], TYPES.get(s['albumType'], s['albumType']),
                 o['kind'], o['album'] or '', o['year'] or '', o['url'] or '', o['how'],
                 s['release'], s['disc'], s['track'], s['ms'], f'=O{i}/86400000', 'Sí' if s['explicit'] else 'No',
                 'Sí' if s['playable'] else 'No', s['isrc'], datetime.strptime(s['added'], '%Y-%m-%dT%H:%M:%SZ'),
                 s['url'], s['id'], s['albumId'], oid(o)])
sheet(ws, ['#', 'Canción', 'Artistas', 'Disco', 'Artistas del disco', 'Tipo de disco',
           'Caso', 'Álbum original', 'Año del original', 'Enlace del original', 'Cómo se encontró', 'Lanzamiento', 'Disco nº', 'Pista nº',
           'Duración (ms)', 'Duración', 'Explícita', 'Reproducible', 'ISRC', 'Añadida a la playlist (UTC)', 'Enlace', 'ID canción', 'ID disco', 'ID del original'],
      [6, 36, 28, 32, 28, 15, 22, 34, 10, 12, 48, 12, 9, 9, 13, 10, 10, 13, 15, 22, 18, 25, 25, 25], rows,
      {'A': f"Orden en la playlist. Fuente: API web de Spotify, playlist 77Yaz4xSfZjJqRTIXo4ZE2 (snapshot {d['snapshot']}), leída el 2026-10-26.",
       'G': 'Columna añadida. De qué tipo de edición viene la canción: Directa (álbum de estudio), Sencillo, Deluxe o edición especial, En vivo, Recopilatorio (incluye bandas sonoras) o Remasterizado. Clasificado por el tipo de disco y por palabras en el nombre del disco y de la canción.',
       'H': 'Columna añadida. Álbum de estudio donde salió la canción por primera vez. Directa: el mismo disco. Buscado en Spotify por la misma grabación (ISRC) y, si no, por título y artista; se toma el álbum más antiguo que no sea en vivo, recopilatorio, remix ni edición especial.',
       'K': 'Columna añadida. Cómo se obtuvo el álbum original, y si falta verificar algo.',
       'L': 'Fecha de lanzamiento del disco tal como la da Spotify (a veces solo año o año-mes).',
       'P': 'Calculada a partir de Duración (ms).'})
for r in range(2, ws.max_row + 1):
    ws[f'P{r}'].number_format = '[m]:ss'
    ws[f'T{r}'].number_format = 'yyyy-mm-dd hh:mm'
    for col in 'UJ':
        c = ws[f'{col}{r}']
        if c.value:
            c.hyperlink = c.value; c.value = 'Abrir'; c.font = LINK
for col in 'GHIJKX':
    ws[f'{col}1'].fill = NEW

wa = wb.create_sheet('Discos')
rows = []
for i, a in enumerate(d['albums'], start=2):
    rows.append([a['name'], a['artists'], TYPES.get(a['type'], a['type']), a['release'], a['totalTracks'],
                 f'=COUNTIF(Canciones!$W$2:$W${len(d["songs"]) + 1},K{i})', a['seconds'], f'=IF(G{i}="","",G{i}/86400)',
                 a['status'], a['image'], a['id'], a['url'], 'En la playlist', f'=COUNTIF(Canciones!$X$2:$X${len(d["songs"]) + 1},K{i})'])
start = len(d['albums']) + 2
for j, a in enumerate(ADD):
    i = start + j
    rows.append([a['name'], a['artists'], 'Álbum', a['release'], None,
                 f'=COUNTIF(Canciones!$W$2:$W${len(d["songs"]) + 1},K{i})', None, f'=IF(G{i}="","",G{i}/86400)',
                 'Por medir: original encontrado', '', a['id'], a['url'], 'Original encontrado',
                 f'=COUNTIF(Canciones!$X$2:$X${len(d["songs"]) + 1},K{i})'])
sheet(wa, ['Disco', 'Artistas', 'Tipo', 'Lanzamiento', 'Pistas del disco', 'Canciones en la playlist', 'Duración (seg)',
           'Duración del disco', 'Estado en Hill-B', 'Portada', 'ID disco', 'Enlace', 'Origen', 'Canciones que tienen este disco como original'],
      [36, 30, 15, 12, 14, 13, 14, 15, 28, 16, 25, 10, 20, 18], rows,
      {'F': 'Cuenta las canciones de este disco en la pestaña Canciones.',
       'G': 'Duración del disco entero (suma de todas sus pistas), medida con la API de Spotify. Vacía si aún no se midió.',
       'M': 'Columna añadida. En la playlist: disco del que viene alguna canción. Original encontrado: álbum de estudio original de una canción que viene de un sencillo, deluxe, en vivo o recopilatorio; no estaba en la lista (ni el mismo disco ni uno con igual nombre y artista).',
       'N': 'Columna añadida. Canciones de la pestaña Canciones cuyo Álbum original es este disco (incluye las directas).',
       'B': 'En los originales encontrados va el artista principal de la canción.',
       'I': 'En la app: medido y en albums.json. Por medir: falta medirlo, Spotify cortó la cuota diaria. Fuera: sencillos y EP, que la app no propone.'})
for col in 'MN':
    wa[f'{col}1'].fill = NEW
for r in range(2, wa.max_row + 1):
    wa[f'H{r}'].number_format = '[h]:mm:ss'
    for col in 'JL':
        c = wa[f'{col}{r}']
        if c.value:
            c.hyperlink = c.value; c.value = 'Ver' if col == 'J' else 'Abrir'; c.font = LINK
wb.save(out)
