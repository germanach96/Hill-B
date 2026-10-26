# The catalog as an Excel workbook for the owner: one sheet of songs, one of albums.
# The owner reviews it and says which originals they found by hand (those go in data/manual.json).
#
#   python3 tools/excel.py "<output.xlsx>"        reads data/catalog.json (needs openpyxl)
import json
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
catalog = json.loads((ROOT / 'data' / 'catalog.json').read_text())
songs, albums = catalog['songs'], catalog['albums']
artist_tags = catalog.get('artistTags', {})

TYPES = {'album': 'Álbum', 'compilation': 'Recopilatorio', 'single': 'Sencillo o EP'}
CASES = {'direct': 'Directa', 'single': 'Sencillo', 'edition': 'Deluxe o edición especial',
         'live': 'En vivo', 'compilation': 'Recopilatorio', 'remaster': 'Remasterizado'}
SOURCES = {'playlist': 'En la playlist', 'original': 'Original encontrado', 'manual': 'Puesto a mano'}
FONT = Font(name='Arial', size=10)
HEAD = Font(name='Arial', size=10, bold=True, color='FFFFFF')
FILL = PatternFill('solid', fgColor='1F2937')
FILL_ORIGINAL = PatternFill('solid', fgColor='1C5CAB')  # columns about the original album
LINK = Font(name='Arial', size=10, color='0563C1', underline='single')


def base_album(name):
    return re.sub(r'\s+-\s+.*$', '', re.sub(r'\s*[([].*?[)\]]', '', name)).strip() or name


def original_cells(s):
    """Original album, its year, its link, its id and how it was found, as the owner reads them."""
    own = albums[s['album']]
    o = s.get('original', 'pending')
    if o and o != 'pending':
        a = albums[o]
        name = base_album(a['name']) if 'edición especial' in s['how'] else a['name']
        return name, a['release'][:4], a['url'], o, s['how']
    if o is None:
        return {
            'single': ('Sin álbum (solo sencillo)', '', '', '', 'No aparece en ningún álbum del artista en Spotify'),
            'edition': (base_album(own['name']), '', '', '', 'Edición estándar del mismo disco; Spotify solo tiene esta edición'),
            'remaster': (base_album(own['name']), '', '', '', 'Edición estándar del mismo disco; Spotify solo tiene esta edición'),
            'compilation': (own['name'], own['release'][:4], own['url'], '', 'No aparece en otro disco: este es donde salió'),
        }.get(s['case'], ('', '', '', '', 'No encontrado'))
    if s['case'] in ('edition', 'remaster'):
        return base_album(own['name']), '', '', '', 'Edición estándar, deducida del nombre; falta verificar en Spotify'
    return '', '', '', '', 'Pendiente: falta buscarlo en Spotify'


def sheet(ws, columns, rows, blue=()):
    """columns: (header, width, note or None)."""
    ws.append([c[0] for c in columns])
    for cell, (_, width, note) in zip(ws[1], columns):
        cell.font, cell.alignment = HEAD, Alignment(vertical='center')
        cell.fill = FILL_ORIGINAL if cell.column_letter in blue else FILL
        ws.column_dimensions[cell.column_letter].width = width
        if note:
            cell.comment = Comment(note, 'Hill-B')
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = FONT
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(columns))}{ws.max_row}'


def tag_line(tags):
    """[["stoner rock", 100], …] → "stoner rock 100 · …"; None when Last.fm was not asked yet."""
    if tags is None:
        return ''
    return ' · '.join(f'{name} {weight}' for name, weight in tags) or '(ninguna)'


def links(ws, col, label):
    for r in range(2, ws.max_row + 1):
        cell = ws[f'{col}{r}']
        if cell.value:
            cell.hyperlink, cell.value, cell.font = cell.value, label, LINK


wb = Workbook()
ws = wb.active
ws.title = 'Canciones'
n = len(songs)
rows = []
for i, (sid, s) in enumerate(songs.items(), start=2):
    a = albums[s['album']]
    o_name, o_year, o_url, o_id, how = original_cells(s)
    added = datetime.strptime(s['added'], '%Y-%m-%dT%H:%M:%SZ') if s['added'] else None
    rows.append([i - 1, s['name'], ', '.join(s['artists']), a['name'], ', '.join(a['artists']), TYPES.get(a['type'], a['type']),
                 CASES[s['case']], o_name, o_year, o_url, how, a['release'], s['disc'], s['track'], s['ms'], f'=O{i}/86400000',
                 'Sí' if s['explicit'] else 'No', s['isrc'], added, s['url'], sid, s['album'], o_id])
sheet(ws, [
    ('#', 6, f"Orden de llegada a la playlist {catalog['playlist']}. Fuente: API web de Spotify, catálogo de Hill-B (data/catalog.json)."),
    ('Canción', 36, None), ('Artistas', 28, None), ('Disco', 32, None), ('Artistas del disco', 28, None), ('Tipo de disco', 15, None),
    ('Caso', 22, 'De qué edición viene la canción: Directa (álbum de estudio), Sencillo, Deluxe o edición especial, En vivo, '
                 'Recopilatorio (incluye bandas sonoras) o Remasterizado. Por el tipo de disco y palabras del nombre.'),
    ('Álbum original', 34, 'Álbum de estudio donde salió la canción. Buscado en Spotify por la misma grabación (ISRC) y, si no, '
                           'por título y artista: el más antiguo que no sea en vivo, recopilatorio, remix ni edición especial.'),
    ('Año del original', 10, None), ('Enlace del original', 12, None),
    ('Cómo se encontró', 48, 'Cómo se obtuvo el álbum original y si falta verificar algo. Los vacíos se resuelven a mano.'),
    ('Lanzamiento', 12, 'Fecha del disco tal como la da Spotify (a veces solo año o año-mes).'),
    ('Disco nº', 9, None), ('Pista nº', 9, None), ('Duración (ms)', 13, None), ('Duración', 10, 'Calculada a partir de Duración (ms).'),
    ('Explícita', 10, None), ('ISRC', 15, None), ('Añadida a la playlist (UTC)', 22, None), ('Enlace', 10, None),
    ('ID canción', 25, None), ('ID disco', 25, None), ('ID del original', 25, None),
], rows, blue='GHIJKW')
for r in range(2, ws.max_row + 1):
    ws[f'P{r}'].number_format = '[m]:ss'
    ws[f'S{r}'].number_format = 'yyyy-mm-dd hh:mm'
links(ws, 'J', 'Abrir')
links(ws, 'T', 'Abrir')

wa = wb.create_sheet('Discos')
artists_of = {}
for s in songs.values():
    for aid in (s['album'], s.get('original')):
        if aid:
            artists_of.setdefault(aid, s['artists'][0])
rows = []
for i, (aid, a) in enumerate(albums.items(), start=2):
    status = 'Fuera: sencillo o EP' if a['type'] == 'single' else 'En la app' if (a['seconds'] or 0) > 0 else 'Por medir'
    rows.append([a['name'], ', '.join(a['artists']) or artists_of.get(aid, ''), TYPES.get(a['type'], a['type']), a['release'],
                 a['totalTracks'], f'=COUNTIF(Canciones!$V$2:$V${n + 1},K{i})', a['seconds'], f'=IF(G{i}="","",G{i}/86400)',
                 status, a['image'] or '', aid, a['url'], SOURCES[a['source']], f'=COUNTIF(Canciones!$W$2:$W${n + 1},K{i})',
                 tag_line(a.get('tags')), tag_line(artist_tags.get(a['artists'][0])) if a['artists'] else ''])
sheet(wa, [
    ('Disco', 36, None), ('Artistas', 30, 'En los originales aún sin medir va el artista principal de la canción.'),
    ('Tipo', 15, None), ('Lanzamiento', 12, None), ('Pistas del disco', 14, None),
    ('Canciones en la playlist', 13, 'Canciones de la pestaña Canciones que vienen de este disco.'),
    ('Duración (seg)', 14, 'Duración del disco entero (suma de todas sus pistas), medida con la API de Spotify. Vacía si aún no se midió.'),
    ('Duración del disco', 15, None),
    ('Estado en Hill-B', 20, 'En la app: medido y en albums.json. Por medir: falta medirlo. Fuera: sencillos y EP, que nunca entran.'),
    ('Portada', 10, None), ('ID disco', 25, None), ('Enlace', 10, None),
    ('Origen', 20, 'En la playlist: disco del que viene alguna canción. Original encontrado: álbum de estudio original de una '
                   'canción de un sencillo, deluxe, en vivo o recopilatorio. Puesto a mano: original que encontró el dueño.'),
    ('Canciones que tienen este disco como original', 18, 'Canciones cuyo Álbum original es este disco (incluye las directas).'),
    ('Etiquetas del disco (Last.fm)', 50, 'Etiquetas que la gente pone al disco en Last.fm, con su peso de 0 a 100 relativo a la más '
                                        'votada. (ninguna): Last.fm no tiene. Vacía: aún sin consultar.'),
    ('Etiquetas del artista (Last.fm)', 50, 'Etiquetas del artista principal en Last.fm, con el mismo peso de 0 a 100.'),
], rows, blue='MN')
for r in range(2, wa.max_row + 1):
    wa[f'H{r}'].number_format = '[h]:mm:ss'
links(wa, 'J', 'Ver')
links(wa, 'L', 'Abrir')

wb.save(sys.argv[1])
print(f'{n} canciones y {len(albums)} discos → {sys.argv[1]}')
